from fastapi import FastAPI, HTTPException, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime, date
import os
import hmac
import hashlib
import base64
import json


# =========================================================
# APP
# =========================================================

app = FastAPI(title="Rahul Software API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

security = HTTPBearer(auto_error=False)

BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "2026-27.xlsm"

USERNAME = os.getenv("RAHUL_USERNAME", "rahul")
PASSWORD = os.getenv("RAHUL_PASSWORD", "rahul123")
SECRET = os.getenv("RAHUL_SECRET", "change-this-secret")


# =========================================================
# HELPERS
# =========================================================

def get_workbook():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    return load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True,
        keep_vba=True
    )


def clean_value(value):
    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return value


def make_token(username):
    payload = {
        "username": username,
        "exp": int(datetime.now().timestamp()) + 86400
    }

    raw = json.dumps(
        payload,
        separators=(",", ":")
    ).encode()

    encoded = base64.urlsafe_b64encode(raw).decode().rstrip("=")

    signature = hmac.new(
        SECRET.encode(),
        encoded.encode(),
        hashlib.sha256
    ).hexdigest()

    return encoded + "." + signature


def verify_token(token):
    try:
        encoded, signature = token.split(".", 1)

        expected = hmac.new(
            SECRET.encode(),
            encoded.encode(),
            hashlib.sha256
        ).hexdigest()

        if not hmac.compare_digest(
            signature,
            expected
        ):
            return False

        padding = "=" * (-len(encoded) % 4)

        raw = base64.urlsafe_b64decode(
            encoded + padding
        )

        payload = json.loads(raw.decode())

        if payload["exp"] < int(
            datetime.now().timestamp()
        ):
            return False

        return True

    except Exception:
        return False


def require_auth(
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    if credentials is None:
        raise HTTPException(
            status_code=401,
            detail="Authentication required"
        )

    if credentials.scheme.lower() != "bearer":
        raise HTTPException(
            status_code=401,
            detail="Invalid authentication"
        )

    if not verify_token(credentials.credentials):
        raise HTTPException(
            status_code=401,
            detail="Invalid or expired token"
        )

    return True


def make_headers(headers):
    result = []

    for i, header in enumerate(headers):
        name = str(header).strip() if header is not None else ""

        if not name:
            name = f"Column {i + 1}"

        original = name
        counter = 2

        while name in result:
            name = f"{original} ({counter})"
            counter += 1

        result.append(name)

    return result


def rows_from_sheet(ws):
    rows = ws.iter_rows(values_only=True)

    try:
        raw_headers = next(rows)
    except StopIteration:
        return [], []

    headers = make_headers(raw_headers)

    data = []

    for row in rows:
        if not any(value is not None for value in row):
            continue

        record = {}

        for i, header in enumerate(headers):
            value = row[i] if i < len(row) else None
            record[header] = clean_value(value)

        data.append(record)

    return headers, data


def number(value):
    try:
        if value is None or value == "":
            return 0

        return float(str(value).replace(",", "").strip())

    except Exception:
        return 0


def date_only(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    for fmt in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(
                text,
                fmt
            ).date()
        except Exception:
            pass

    return None


def text(value):
    if value is None:
        return ""

    return str(value).strip()


def first_existing(record, names):
    for name in names:
        if name in record:
            return record[name]

    return None


# =========================================================
# HOME
# =========================================================

@app.get("/")
def home():
    return {
        "status": "online",
        "message": "Rahul Software API is running"
    }


# =========================================================
# LOGIN
# =========================================================

@app.post("/login")
def login(data: dict):

    username = text(data.get("username"))
    password = text(data.get("password"))

    if (
        hmac.compare_digest(username, USERNAME)
        and hmac.compare_digest(password, PASSWORD)
    ):
        token = make_token(username)

        return {
            "status": "success",
            "access_token": token,
            "token": token,
            "token_type": "bearer"
        }

    raise HTTPException(
        status_code=401,
        detail="Invalid username or password"
    )


# =========================================================
# SHEETS
# =========================================================

@app.get("/sheets")
def sheets(authenticated: bool = Depends(require_auth)):

    try:
        wb = get_workbook()

        names = wb.sheetnames

        wb.close()

        return {
            "count": len(names),
            "sheets": names
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =========================================================
# SINGLE SHEET
# =========================================================

@app.get("/sheet/{sheet_name}")
def get_sheet(
    sheet_name: str,
    limit: int = 100,
    authenticated: bool = Depends(require_auth)
):

    try:
        wb = get_workbook()

        if sheet_name not in wb.sheetnames:
            wb.close()

            raise HTTPException(
                status_code=404,
                detail=f"Sheet '{sheet_name}' not found"
            )

        ws = wb[sheet_name]

        headers, data = rows_from_sheet(ws)

        wb.close()

        data = data[:max(1, min(limit, 1000))]

        return {
            "sheet": sheet_name,
            "headers": headers,
            "count": len(data),
            "data": data
        }

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =========================================================
# GLOBAL SEARCH
# =========================================================

@app.get("/search")
def search(
    q: str,
    sheet: str = "ALL SHEETS",
    authenticated: bool = Depends(require_auth)
):

    try:
        wb = get_workbook()

        search_text = q.strip().lower()

        if not search_text:
            wb.close()

            return {
                "query": q,
                "sheet": sheet,
                "count": 0,
                "results": []
            }

        results = []

        for sheet_name in wb.sheetnames:

            if (
                sheet != "ALL SHEETS"
                and sheet_name != sheet
            ):
                continue

            ws = wb[sheet_name]

            headers, data = rows_from_sheet(ws)

            for record in data:

                found = any(
                    search_text in text(value).lower()
                    for value in record.values()
                )

                if found:

                    row = [
                        record.get(header)
                        for header in headers
                    ]

                    results.append({
                        "sheet": sheet_name,
                        "headers": headers,
                        "row": row,
                        "data": record
                    })

        wb.close()

        return {
            "query": q,
            "sheet": sheet,
            "count": len(results),
            "results": results
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =========================================================
# DASHBOARD
# =========================================================

@app.get("/dashboard")
def dashboard(
    authenticated: bool = Depends(require_auth)
):

    try:
        wb = get_workbook()

        today = datetime.now().date()

        # -------------------------------------------------
        # ORDERS
        # -------------------------------------------------

        orders = []

        if "ORDERS" in wb.sheetnames:
            _, orders = rows_from_sheet(
                wb["ORDERS"]
            )

        order_quantity = 0
        order_numbers = set()

        today_orders = 0
        today_order_quantity = 0

        size_data = {}
        party_data = {}

        for row in orders:

            order_no = first_existing(
                row,
                ["Order No.", "Order No"]
            )

            quantity = number(
                first_existing(
                    row,
                    [
                        "Quantity Pcs.",
                        "Quantity",
                        "QUANTITY"
                    ]
                )
            )

            row_date = date_only(
                first_existing(
                    row,
                    ["Date", "DATE"]
                )
            )

            size = text(
                first_existing(
                    row,
                    ["Size", "SIZE"]
                )
            )

            party = text(
                first_existing(
                    row,
                    [
                        "PARTY NAME",
                        "Party Name",
                        "PARTY"
                    ]
                )
            )

            if order_no not in (
                None,
                ""
            ):
                order_numbers.add(
                    text(order_no)
                )

            order_quantity += quantity

            if row_date == today:
                today_orders += 1
                today_order_quantity += quantity

            if size:
                if size not in size_data:
                    size_data[size] = {
                        "size": size,
                        "orders": 0,
                        "quantity": 0
                    }

                size_data[size]["orders"] += 1
                size_data[size]["quantity"] += quantity

            if party:
                if party not in party_data:
                    party_data[party] = {
                        "party": party,
                        "orders": 0,
                        "quantity": 0
                    }

                party_data[party]["orders"] += 1
                party_data[party]["quantity"] += quantity

        # -------------------------------------------------
        # DISPATCH
        # -------------------------------------------------

        dispatch = []

        if "Dispatched Orders" in wb.sheetnames:
            _, dispatch = rows_from_sheet(
                wb["Dispatched Orders"]
            )

        dispatch_quantity = 0
        today_dispatch = 0
        today_dispatch_quantity = 0

        for row in dispatch:

            quantity = number(
                first_existing(
                    row,
                    [
                        "Quantity Pcs.",
                        "Quantity",
                        "QUANTITY"
                    ]
                )
            )

            dispatch_quantity += quantity

            row_date = date_only(
                first_existing(
                    row,
                    [
                        "Date",
                        "BILL DATE",
                        "IN/OUT DATE"
                    ]
                )
            )

            if row_date == today:
                today_dispatch += 1
                today_dispatch_quantity += quantity

        # -------------------------------------------------
        # HOLD ORDERS
        # -------------------------------------------------

        hold = []

        if "Holded Orders" in wb.sheetnames:
            _, hold = rows_from_sheet(
                wb["Holded Orders"]
            )

        hold_quantity = 0

        for row in hold:

            hold_quantity += number(
                first_existing(
                    row,
                    [
                        "Quantity Pcs.",
                        "Quantity",
                        "QUANTITY"
                    ]
                )
            )

        wb.close()

        # -------------------------------------------------
        # CLEAN NUMBERS
        # -------------------------------------------------

        def clean_number(value):

            if float(value).is_integer():
                return int(value)

            return round(value, 2)

        size_wise = sorted(
            [
                {
                    "size": item["size"],
                    "orders": item["orders"],
                    "quantity": clean_number(
                        item["quantity"]
                    )
                }
                for item in size_data.values()
            ],
            key=lambda x: x["quantity"],
            reverse=True
        )

        party_wise = sorted(
            [
                {
                    "party": item["party"],
                    "orders": item["orders"],
                    "quantity": clean_number(
                        item["quantity"]
                    )
                }
                for item in party_data.values()
            ],
            key=lambda x: x["quantity"],
            reverse=True
        )

        return {
            "date": today.isoformat(),

            "today": {
                "orders": today_orders,
                "order_quantity": clean_number(
                    today_order_quantity
                ),
                "dispatch": today_dispatch,
                "dispatch_quantity": clean_number(
                    today_dispatch_quantity
                )
            },

            "orders": {
                "total_rows": len(orders),
                "unique_orders": len(order_numbers),
                "total_quantity": clean_number(
                    order_quantity
                )
            },

            "dispatch": {
                "total_rows": len(dispatch),
                "total_quantity": clean_number(
                    dispatch_quantity
                ),
                "today_rows": today_dispatch,
                "today_quantity": clean_number(
                    today_dispatch_quantity
                )
            },

            "hold": {
                "orders": len(hold),
                "quantity": clean_number(
                    hold_quantity
                )
            },

            "size_wise": size_wise,
            "party_wise": party_wise,

            "last_updated": datetime.now().astimezone().isoformat()
        }

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )
